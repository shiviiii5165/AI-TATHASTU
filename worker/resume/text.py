"""Stage 2 — text and layout extraction into a Document with char offsets.

Every block keeps its page and bbox, so any span of Document.text can be
mapped back to where it appeared (profile_facts.location).
"""

from __future__ import annotations

import io
import shutil
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

# 200 dpi grayscale keeps Tesseract accurate on resume-sized text while staying
# well inside a 512 MB instance (300 dpi colour crashed the free Render worker).
OCR_DPI = 200
MAX_OCR_PAGES = 6
OCR_UNCERTAIN_BELOW = 0.70
# A PDF page with less text than this is treated as scanned.
MIN_TEXT_CHARS = 40
# Image resumes: refuse decompression bombs, downscale huge photos before OCR.
MAX_IMAGE_PIXELS = 40_000_000
MAX_IMAGE_SIDE = 3500
# PDFs: a 10 MB file can hold thousands of pages, or pages far larger than paper.
MAX_PDF_PAGES = 50
MAX_PDF_PAGE_POINTS = 2000  # about 28 inches; A0 is 3370, A3 is 1191

BBox = Tuple[float, float, float, float]


class OcrUnavailable(Exception):
    """The file needs OCR but the tesseract binary isn't installed."""


@dataclass
class Block:
    text: str
    page: Optional[int]
    bbox: Optional[BBox]
    start: int = 0
    end: int = 0


@dataclass
class PageInfo:
    number: int
    method: str  # "text" | "ocr" | "skipped" (past MAX_OCR_PAGES)
    ocr_confidence: Optional[float] = None

    @property
    def uncertain(self) -> bool:
        return self.method == "ocr" and (self.ocr_confidence or 0) < OCR_UNCERTAIN_BELOW


@dataclass
class Document:
    blocks: List[Block]
    pages: List[PageInfo] = field(default_factory=list)
    text: str = ""
    # PDF text blocks that extracted as whitespace only: fonts without a
    # Unicode map (often a designed header with name/contact) read as nothing.
    blank_blocks: int = 0

    def __post_init__(self) -> None:
        parts, pos = [], 0
        for b in self.blocks:
            b.start, b.end = pos, pos + len(b.text)
            parts.append(b.text)
            pos = b.end + 1  # "\n" separator
        self.text = "\n".join(parts)

    def locate(self, start: int, end: int) -> Optional[Dict[str, object]]:
        if start < 0:  # vision fact with no matching text span
            return None
        for b in self.blocks:
            if b.start <= start <= b.end:  # the block where the span begins
                if b.page is None:
                    return None
                return {"page": b.page, "bbox": list(b.bbox) if b.bbox else [0, 0, 0, 0]}
        return None

    def page_info(self, start: int) -> Optional[PageInfo]:
        if start < 0:
            return None
        loc = self.locate(start, start + 1)
        if not loc:
            return None
        return next((p for p in self.pages if p.number == loc["page"]), None)

    @property
    def blank_ratio(self) -> float:
        return self.blank_blocks / max(self.blank_blocks + len(self.blocks), 1)

    @property
    def used_ocr(self) -> bool:
        return any(p.method == "ocr" for p in self.pages)


def extract(data: bytes, kind: str, force_ocr: bool = False) -> Document:
    """force_ocr (PDF only): ignore the text layer and OCR every page, for PDFs
    whose embedded text is garbage (see resume/quality.py)."""
    if kind == "pdf":
        return _pdf(data, force_ocr)
    if kind == "docx":
        return _docx(data)
    if kind == "image":
        return _image(data)
    return _txt(data)


# ─── PDF ────────────────────────────────────────────────────────────────────


def _pdf(data: bytes, force_ocr: bool = False) -> Document:
    import fitz  # PyMuPDF

    blocks: List[Block] = []
    pages: List[PageInfo] = []
    blank = 0
    with fitz.open(stream=data, filetype="pdf") as pdf:
        if pdf.page_count > MAX_PDF_PAGES:
            raise ValueError(f"PDF has {pdf.page_count} pages (limit {MAX_PDF_PAGES})")
        for index, page in enumerate(pdf):
            if max(page.rect.width, page.rect.height) > MAX_PDF_PAGE_POINTS:
                raise ValueError("PDF page is too large to process")  # its bitmap wouldn't fit in memory
            number = index + 1
            text_blocks = [b for b in page.get_text("blocks", sort=True) if b[6] == 0]
            raw = [b for b in text_blocks if b[4].strip()]
            blank += len(text_blocks) - len(raw)
            if not force_ocr and sum(len(b[4].strip()) for b in raw) >= MIN_TEXT_CHARS:
                pages.append(PageInfo(number, "text"))
                for x0, y0, x1, y1, text, *_ in _reading_order(raw, page.rect.width):
                    blocks.append(Block(_clean(text), number, (x0, y0, x1, y1)))
            elif index >= MAX_OCR_PAGES:
                pages.append(PageInfo(number, "skipped"))
            else:
                pix = page.get_pixmap(dpi=OCR_DPI, colorspace=fitz.csGRAY)
                ocr_blocks, conf = _ocr_png(pix.tobytes("png"), number, scale=72 / OCR_DPI)
                del pix  # free the bitmap before the next page
                pages.append(PageInfo(number, "ocr", conf))
                blocks.extend(ocr_blocks)
    return Document(blocks, pages, blank_blocks=0 if force_ocr else blank)


def _reading_order(raw: list, width: float) -> list:
    """
    Two-column layouts: read the full-width header, then the left column top
    to bottom, then the right. Single-column pages keep PyMuPDF's sort.
    """
    mid = width / 2
    left = [b for b in raw if b[2] <= mid + width * 0.05]
    right = [b for b in raw if b[0] >= mid - width * 0.05 and b not in left]
    if len(left) < 3 or len(right) < 3:
        return raw
    full = [b for b in raw if b not in left and b not in right]
    top = min(b[1] for b in left + right)
    header = [b for b in full if b[3] <= top]
    rest = [b for b in full if b not in header]
    by_y = lambda bs: sorted(bs, key=lambda b: (b[1], b[0]))  # noqa: E731
    return by_y(header) + by_y(left) + by_y(right) + by_y(rest)


def _clean(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


# ─── OCR ────────────────────────────────────────────────────────────────────


def _require_tesseract() -> None:
    if not shutil.which("tesseract"):
        raise OcrUnavailable("tesseract is not installed on this worker")


def _ocr_png(png: bytes, page: int, scale: float = 1.0) -> Tuple[List[Block], float]:
    _require_tesseract()
    import pytesseract
    from PIL import Image

    image = Image.open(io.BytesIO(png))
    data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
    groups: Dict[Tuple[int, int], Dict[str, object]] = {}
    confs: List[float] = []
    for i, word in enumerate(data["text"]):
        word = (word or "").strip()
        conf = float(data["conf"][i])
        if not word or conf < 0:
            continue
        confs.append(conf)
        key = (data["block_num"][i], data["par_num"][i])
        g = groups.setdefault(key, {"lines": {}, "box": [1e9, 1e9, 0.0, 0.0]})
        g["lines"].setdefault(data["line_num"][i], []).append(word)  # type: ignore[union-attr]
        x, y, w, h = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
        box = g["box"]  # type: ignore[assignment]
        box[0], box[1] = min(box[0], x), min(box[1], y)
        box[2], box[3] = max(box[2], x + w), max(box[3], y + h)

    blocks = []
    for key in sorted(groups):
        g = groups[key]
        text = "\n".join(" ".join(ws) for _, ws in sorted(g["lines"].items()))  # type: ignore[union-attr]
        x0, y0, x1, y1 = (v * scale for v in g["box"])  # type: ignore[union-attr]
        blocks.append(Block(text, page, (x0, y0, x1, y1)))
    confidence = (sum(confs) / len(confs) / 100) if confs else 0.0
    return blocks, confidence


def _image(data: bytes) -> Document:
    from PIL import Image, ImageOps

    _require_tesseract()
    image = Image.open(io.BytesIO(data))
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise ValueError("image is too large to OCR")
    # Phone photos store rotation in EXIF; Tesseract reads the raw pixels.
    image = ImageOps.exif_transpose(image).convert("L")
    if max(image.size) > MAX_IMAGE_SIDE:  # bounded memory; still ~300 dpi for A4
        image.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    blocks, conf = _ocr_png(buf.getvalue(), 1)
    return Document(blocks, [PageInfo(1, "ocr", conf)])


# ─── DOCX / TXT ─────────────────────────────────────────────────────────────


def _docx(data: bytes) -> Document:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    blocks: List[Block] = []
    # Many templates put the name and contact line in the page header; read it
    # first (once, even when sections repeat it), as it appears above the body.
    seen_header: set = set()
    for section in document.sections:
        for para in section.header.paragraphs:
            text = _clean(para.text)
            if text and text not in seen_header:
                seen_header.add(text)
                blocks.append(Block(text, None, None))
    # Walk the body in order so tables stay where they appear.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _clean(Paragraph(child, document).text)
            if text:
                blocks.append(Block(text, None, None))
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = []
                for cell in row.cells:
                    t = _clean(cell.text)
                    if t and t not in cells:  # merged cells repeat
                        cells.append(t)
                if cells:
                    blocks.append(Block(" | ".join(cells), None, None))
    # DOCX has no fixed pages; facts get no page number.
    return Document(blocks, [])


def _txt(data: bytes) -> Document:
    try:
        text = data.decode("utf-8-sig")  # drops a BOM, which would otherwise glue onto the name
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    paras = [_clean(p) for p in text.replace("\r\n", "\n").split("\n\n")]
    return Document([Block(p, None, None) for p in paras if p], [])
